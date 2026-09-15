#!/usr/bin/env python3
"""Build the inventory workbook, including optional controlled-crawl evidence.

The original Summary, Domains, Subdomains, and Web Applications sheets are
always generated. The three additional sheets are empty-but-labelled when a
crawl was not requested, preserving the lightweight inventory workflow.
"""
from __future__ import annotations

import argparse
import json
import os
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SEVERITY_FILLS = {
    "Critical": PatternFill(start_color="C00000", end_color="C00000", fill_type="solid"),
    "High": PatternFill(start_color="ED7D31", end_color="ED7D31", fill_type="solid"),
    "Medium": PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid"),
    "Low": PatternFill(start_color="A9D18E", end_color="A9D18E", fill_type="solid"),
    "Informational": PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid"),
}


def style_header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")


def autosize(ws, min_width: int = 10, max_width: int = 70) -> None:
    for col_cells in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = max(min_width, min(max_width, length + 2))


def read_lines(path: str | None) -> list[str]:
    if not path or not os.path.exists(path):
        return []
    return sorted({line.strip() for line in Path(path).read_text(encoding="utf-8").splitlines() if line.strip()})


def read_jsonl(path: str | None) -> list[dict]:
    records = []
    if not path or not os.path.exists(path):
        return records
    for line in Path(path).read_text(encoding="utf-8").splitlines():
        try:
            records.append(json.loads(line))
        except json.JSONDecodeError:
            continue
    return records


def read_json(path: str | None) -> dict:
    if not path or not os.path.exists(path):
        return {}
    try:
        return json.loads(Path(path).read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return {}


def read_httpx_json(path: str | None) -> list[dict]:
    return read_jsonl(path)


def add_hyperlink(cell, value: str) -> None:
    if value and value.startswith(("http://", "https://")):
        cell.hyperlink = value
        cell.style = "Hyperlink"


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the asset inventory Excel workbook")
    parser.add_argument("--domains", required=True)
    parser.add_argument("--wildcards", default=None)
    parser.add_argument("--subdomains", required=True)
    parser.add_argument("--httpx-json", required=True)
    parser.add_argument("--scope-normalized", default=None)
    parser.add_argument("--js-assets", default=None)
    parser.add_argument("--secret-findings", default=None)
    parser.add_argument("--crawl-stats", default=None)
    parser.add_argument("--analysis-stats", default=None)
    parser.add_argument("--output", default="asset_inventory.xlsx")
    args = parser.parse_args()

    domains = read_lines(args.domains)
    wildcards = set(read_lines(args.wildcards))
    subdomains = read_lines(args.subdomains)
    httpx_records = read_httpx_json(args.httpx_json)
    normalized = read_json(args.scope_normalized).get("records", [])
    js_assets = read_jsonl(args.js_assets)
    findings = read_jsonl(args.secret_findings)
    crawl_stats, analysis_stats = read_json(args.crawl_stats), read_json(args.analysis_stats)

    scope_by_host: dict[str, list[dict]] = {}
    for record in normalized:
        if record.get("web_target"):
            scope_by_host.setdefault(record.get("host", ""), []).append(record)

    workbook = Workbook()

    # Summary remains first, as in the existing report.
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Asset Inventory Summary"])
    summary["A1"].font = Font(size=14, bold=True)
    summary.append([])
    summary.append(["Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    summary.append(["Domains in scope", len(domains)])
    summary.append(["Unique subdomains discovered", len(subdomains)])

    seen_web_apps = set()
    for record in httpx_records:
        seen_web_apps.add((record.get("url", ""), record.get("status_code", "")))
    summary.append(["Live web applications (httpx)", len(seen_web_apps)])
    summary.append(["Pages visited (optional crawl)", crawl_stats.get("pages_visited", 0)])
    summary.append(["JavaScript assets analyzed (optional)", analysis_stats.get("javascript_assets_analyzed", 0)])
    summary.append(["Secret findings (unverified)", len(findings)])
    summary.append([])
    summary.append(["Sheet", "Contents"])
    summary.append(["Domains", "Normalized in-scope web domains and wildcard roots"])
    summary.append(["Subdomains", "Passive subdomain-enumeration results"])
    summary.append(["Web Applications", "HTTP(S) applications reported by httpx"])
    summary.append(["JavaScript Assets", "Public JS resources discovered during the optional controlled crawl"])
    summary.append(["Secret Findings", "Safely redacted, unverified pattern matches only"])
    summary.append(["Crawl Statistics", "Crawl limits, discovery counts, skips, errors, and analysis totals"])
    for row in range(3, 10):
        summary[f"A{row}"].font = Font(bold=True)
    style_header(summary, row=11)
    autosize(summary)

    domains_sheet = workbook.create_sheet("Domains")
    domains_sheet.append(["Domain", "Type", "Notes"])
    for domain in domains:
        records = scope_by_host.get(domain, [])
        notes = " | ".join(dict.fromkeys(record.get("instructions", "") for record in records if record.get("instructions")))
        domain_type = "Wildcard root" if domain in wildcards else "Explicit domain"
        domains_sheet.append([domain, domain_type, notes])
    style_header(domains_sheet)
    domains_sheet.freeze_panes = "A2"
    autosize(domains_sheet)

    subdomains_sheet = workbook.create_sheet("Subdomains")
    subdomains_sheet.append(["Subdomain", "Root Domain (best guess)"])
    for subdomain in subdomains:
        candidates = [domain for domain in domains if subdomain == domain or subdomain.endswith("." + domain)]
        subdomains_sheet.append([subdomain, max(candidates, key=len) if candidates else ""])
    style_header(subdomains_sheet)
    subdomains_sheet.freeze_panes = "A2"
    autosize(subdomains_sheet)

    web_sheet = workbook.create_sheet("Web Applications")
    web_sheet.append(["URL", "Host", "Status Code", "Title", "Technologies", "Content Length", "Redirect Location", "Final URL"])
    seen = set()
    for record in httpx_records:
        url = record.get("url", "")
        status = record.get("status_code", "")
        key = (url, status)
        if key in seen:
            continue
        seen.add(key)
        row = [url, record.get("host", record.get("input", "")), status, record.get("title", ""),
               ", ".join(record.get("tech", []) or []), record.get("content_length", ""),
               record.get("location", ""), record.get("final_url", url)]
        web_sheet.append(row)
        add_hyperlink(web_sheet.cell(web_sheet.max_row, 1), url)
        add_hyperlink(web_sheet.cell(web_sheet.max_row, 8), row[7])
    style_header(web_sheet)
    web_sheet.freeze_panes = "A2"
    autosize(web_sheet)

    js_sheet = workbook.create_sheet("JavaScript Assets")
    js_headers = ["URL", "Host", "Parent Page", "Status Code", "Content Type", "Content Length", "Hash", "Retrieved At", "Source Map", "In Scope"]
    js_sheet.append(js_headers)
    for asset in js_assets:
        row = [asset.get("url", ""), asset.get("host", ""), asset.get("parent_page", ""), asset.get("status_code", ""),
               asset.get("content_type", ""), asset.get("content_length", ""), asset.get("hash", ""), asset.get("retrieved_at", ""),
               asset.get("source_map", ""), asset.get("in_scope", False)]
        js_sheet.append(row)
        add_hyperlink(js_sheet.cell(js_sheet.max_row, 1), row[0])
        add_hyperlink(js_sheet.cell(js_sheet.max_row, 3), row[2])
        add_hyperlink(js_sheet.cell(js_sheet.max_row, 9), row[8])
    style_header(js_sheet)
    js_sheet.freeze_panes = "A2"
    autosize(js_sheet)

    finding_sheet = workbook.create_sheet("Secret Findings")
    finding_headers = ["Finding ID", "Host", "Source URL", "Original JS URL", "Finding Type", "Severity", "Confidence", "Matched Pattern", "Redacted Evidence", "Context", "Line Number", "First Seen", "Verification Status"]
    finding_sheet.append(finding_headers)
    severity_column = finding_headers.index("Severity") + 1
    for finding in findings:
        row = [finding.get("finding_id", ""), finding.get("host", ""), finding.get("source_url", ""), finding.get("original_js_url", ""),
               finding.get("finding_type", ""), finding.get("severity", ""), finding.get("confidence", ""), finding.get("matched_pattern", ""),
               finding.get("redacted_evidence", ""), finding.get("context", ""), finding.get("line_number", ""), finding.get("first_seen", ""), finding.get("status", "")]
        finding_sheet.append(row)
        add_hyperlink(finding_sheet.cell(finding_sheet.max_row, 3), row[2])
        add_hyperlink(finding_sheet.cell(finding_sheet.max_row, 4), row[3])
        severity = str(row[5])
        if severity in SEVERITY_FILLS:
            finding_sheet.cell(finding_sheet.max_row, severity_column).fill = SEVERITY_FILLS[severity]
    style_header(finding_sheet)
    finding_sheet.freeze_panes = "A2"
    finding_sheet.auto_filter.ref = f"A1:M{finding_sheet.max_row}"
    autosize(finding_sheet)

    stats_sheet = workbook.create_sheet("Crawl Statistics")
    stats_sheet.append(["Metric", "Value"])
    metrics = {
        "Hosts crawled": crawl_stats.get("hosts_crawled", 0),
        "Pages visited": crawl_stats.get("pages_visited", 0),
        "JS assets discovered": crawl_stats.get("javascript_assets_discovered", 0),
        "JS assets analyzed": analysis_stats.get("javascript_assets_analyzed", 0),
        "Source maps discovered": analysis_stats.get("source_maps_discovered", 0),
        "Source maps analyzed": analysis_stats.get("source_maps_analyzed", 0),
        "Errors/timeouts": crawl_stats.get("errors_timeouts", 0) + analysis_stats.get("analysis_errors_timeouts", 0),
        "Skipped/out-of-scope URLs": crawl_stats.get("skipped_out_of_scope_urls", 0),
        "Crawler requests": crawl_stats.get("requests", 0),
        "Analyzer requests": analysis_stats.get("requests", 0),
    }
    for metric, value in metrics.items():
        stats_sheet.append([metric, value])
    for severity, count in sorted(analysis_stats.get("findings_by_severity", {}).items()):
        stats_sheet.append([f"Findings – {severity}", count])
    for finding_type, count in sorted(analysis_stats.get("findings_by_type", {}).items()):
        stats_sheet.append([f"Findings type – {finding_type}", count])
    for reason, count in sorted(crawl_stats.get("skipped_urls", {}).items()):
        stats_sheet.append([f"Skipped – {reason}", count])
    style_header(stats_sheet)
    stats_sheet.freeze_panes = "A2"
    autosize(stats_sheet)

    workbook.save(args.output)
    print(f"[+] Workbook written to {args.output}")
    print(f"    Domains: {len(domains)} | Subdomains: {len(subdomains)} | Web apps: {len(seen)} | Findings: {len(findings)}")


if __name__ == "__main__":
    main()
